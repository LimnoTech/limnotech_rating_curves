import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import pandas as pd

from .. import settings
from ..model_selection.metrics import fit_metrics

log = logging.getLogger(__name__)

#: Characters a translated formula may contain. Anything else means the formula uses
#: a construct this module does not model, and it is refused rather than evaluated.
_SAFE_EXPRESSION = re.compile(r"^[0-9xX.+\-*/()eE, ]*$")


@dataclass
class WorkbookRating:
    """The rating one field workbook carries, as a curve on the NAVD88 stage axis.

    Attributes
    ----------
    sensor : str
        Sensor id, e.g. ``"SBR-09"``.
    path : pathlib.Path
        The workbook.
    form : str
        ``'linear'``, ``'quadratic'`` or ``'exponential'``.
    formula : str
        The workbook's own Excel formula, verbatim.
    expression : str
        That formula translated to an expression in ``x``, with the coefficient
        cells substituted.
    variable : str
        What ``x`` is: ``'distance_mm'`` or ``'stage_ft'``.
    coefficients : dict
        ``{label: value}`` as the summary sheet labels them (``X2``, ``X``, ``+``).
    radar_elevation_ft : float or None
        The elevation that maps a NAVD88 stage onto the distance axis. None for a
        workbook whose equation is already on the stage axis.
    datum_spread_ft : float or None
        How much that elevation varies across the workbook's visits.
    verified : bool
        Whether re-evaluating the expression reproduced the workbook's own Estimated
        Discharge column.
    note : str
        Why the rating is unusable, when it is.
    """

    sensor: str
    path: Path
    form: str = ""
    formula: str = ""
    expression: str = ""
    variable: str = ""
    coefficients: dict = field(default_factory=dict)
    radar_elevation_ft: float | None = None
    datum_spread_ft: float | None = None
    verified: bool = False
    note: str = ""

    @property
    def ok(self) -> bool:
        """True when this rating can be drawn on the stage axis."""
        return bool(self.expression) and not self.note

    def discharge(self, x) -> np.ndarray:
        """Evaluate the workbook's equation on its own axis.

        Parameters
        ----------
        x : array-like
            Distance to water in mm, or NAVD88 stage in feet - whichever
            ``variable`` says.

        Returns
        -------
        numpy.ndarray
        """
        values = np.asarray(x, dtype=float)
        return np.asarray(eval(self.expression,  # noqa: S307 - allowlisted above
                               {"__builtins__": {}, "exp": np.exp},
                               {"x": values}), dtype=float)

    def discharge_at_stage(self, stage_ft) -> np.ndarray:
        """Evaluate the equation against NAVD88 stage, converting the axis if needed.

        Parameters
        ----------
        stage_ft : array-like
            Water-surface elevation, NAVD88 feet.

        Returns
        -------
        numpy.ndarray
            Discharge in cfs. All-NaN when the workbook has no usable datum tie, so
            a missing conversion shows up as a missing curve rather than as a curve
            in the wrong place.

        Raises
        ------
        ValueError
            If the rating is not usable at all.
        """
        if not self.ok:
            raise ValueError(f"{self.sensor}: {self.note or 'no stored equation'}")
        stage = np.asarray(stage_ft, dtype=float)
        if self.variable == "stage_ft":
            return self.discharge(stage)
        if self.radar_elevation_ft is None:
            return np.full(stage.shape, np.nan)
        return self.discharge((self.radar_elevation_ft - stage) * settings.MM_PER_FOOT)

    def equation(self, precision: int = 6) -> str:
        """The equation as readable text, on its own axis.

        The coefficients are the workbook's, reformatted - substituting them into
        the expression leaves float repr artifacts (``6.000000000000001e-05`` for a
        cell holding ``6e-05``) that would read as spurious precision.
        """
        variable = "d" if self.variable == "distance_mm" else "h"
        text = re.sub(r"\d+\.?\d*(?:[eE][-+]?\d+)?",
                      lambda match: f"{float(match.group()):.{precision}g}",
                      self.expression)
        text = re.sub(r"\bx\b", variable, text)
        return f"Q = {text}".replace("**2", "²").replace("*", "·")

    def curve(self, stage_ft, points: int = settings.GRID_POINTS) -> pd.DataFrame:
        """The rating as a stage / discharge frame over a span of stage.

        Parameters
        ----------
        stage_ft : array-like
            Stage values whose range the curve should span.
        points : int
            Grid resolution.

        Returns
        -------
        pandas.DataFrame
            ``stage_ft`` and ``discharge_cfs``, with non-positive discharge dropped
            - these equations are unconstrained polynomials and several of them go
            negative below the measurements.
        """
        stage = np.asarray(stage_ft, dtype=float)
        low, high = float(np.nanmin(stage)), float(np.nanmax(stage))
        pad = settings.GRID_PAD_FRACTION * ((high - low) or 1.0)
        grid = np.linspace(low - pad, high + pad, points)
        discharge = self.discharge_at_stage(grid)
        frame = pd.DataFrame({"stage_ft": grid, "discharge_cfs": discharge})
        return frame[np.isfinite(frame["discharge_cfs"])
                     & (frame["discharge_cfs"] > 0)].reset_index(drop=True)


# ---------------------------------------------------------------------------
# finding and reading the workbooks
# ---------------------------------------------------------------------------

def workbook_paths(root=None) -> dict:
    """Every current field workbook, keyed by sensor.

    Parameters
    ----------
    root : path-like, optional
        The ``magl_curve_data`` tree. Defaults to ``settings.MAGL_WORKBOOK_ROOT``.

    Returns
    -------
    dict
        ``{sensor: path}``, sorted. Excel lock files (``~$…``) and the superseded
        copies under ``level_data_before_…`` are left out.
    """
    root = Path(settings.MAGL_WORKBOOK_ROOT if root is None else root)
    found = {}
    for path in sorted(root.glob("**/flow@*.xlsx")):
        if path.name.startswith("~$") or "level_data_before" in str(path):
            continue
        match = re.match(r"flow@([A-Za-z]+-\d+)", path.name)
        if match:
            found.setdefault(match.group(1), path)
    return dict(sorted(found.items()))


def _grid(path, sheet, rows, data_only) -> dict:
    """``{coordinate: value}`` over the first `rows` rows of one sheet."""
    import openpyxl
    from openpyxl.utils import get_column_letter
    book = openpyxl.load_workbook(path, data_only=data_only, read_only=True)
    if sheet not in book.sheetnames:
        book.close()
        return {}
    cells = {f"{get_column_letter(cell.column)}{cell.row}": cell.value
             for row in book[sheet].iter_rows(min_row=1, max_row=rows)
             for cell in row if cell.value is not None}
    book.close()
    return cells


def _sheet_names(path) -> list:
    """A workbook's sheet names, without parsing any sheet body."""
    import openpyxl
    book = openpyxl.load_workbook(path, read_only=True)
    names = list(book.sheetnames)
    book.close()
    return names


def _split(coordinate) -> tuple:
    """``'J14'`` -> ``('J', 14)``."""
    match = re.match(r"([A-Z]+)(\d+)$", coordinate)
    return match.group(1), int(match.group(2))


def summary_sheet_name(path, names=None) -> "str | None":
    """The sheet holding the gauging summary and the stored coefficients.

    Named ``Discharge MMT & Elev. Summary``, ``Discharge MMT and Elev. Summary``,
    ``Stage-Discharge`` or ``Stage-discharge`` depending on the workbook, so it is
    identified by its content - the control-point banner in ``A1`` - rather than by
    name. SR-07 spells that banner ``Conrol Point Location=``.
    """
    names = _sheet_names(path) if names is None else names
    likely = [name for name in names if re.search(r"summary|stage.?discharge", name,
                                                  re.I)]
    for name in likely + [name for name in names if name not in likely]:
        banner = _grid(path, name, 1, data_only=True).get("A1")
        if isinstance(banner, str) and "point location" in banner.lower():
            return name
    return None


def _header(cells) -> tuple:
    """``(row, {column letter: header})`` for the summary sheet's Date/Time header."""
    for coordinate, value in cells.items():
        if isinstance(value, str) and value.strip().lower() == "date/time":
            _, row = _split(coordinate)
            return row, {letter: str(text).strip()
                         for letter, text in
                         ((_split(other)[0], text) for other, text in cells.items()
                          if _split(other)[1] == row)}
    return None, {}


def _column_for(headers, *needles) -> "str | None":
    """The column letter whose header contains every needle."""
    for letter, name in headers.items():
        if all(needle in name.lower() for needle in needles):
            return letter
    return None


def _first_formula(cells, letter, header_row, limit=12) -> tuple:
    """``(row, text)`` of the first formula in a column below the header."""
    if letter is None:
        return None, None
    for row in range(header_row + 1, header_row + 1 + limit):
        value = cells.get(f"{letter}{row}")
        if isinstance(value, str) and value.startswith("="):
            return row, value
    return None, None


def _visits(values, headers, header_row, limit=12) -> pd.DataFrame:
    """One row per gauging visit: its date and every summary column's value."""
    rows = [{"date": pd.Timestamp(values[f"A{row}"]).normalize(),
             **{letter: values.get(f"{letter}{row}") for letter in headers}}
            for row in range(header_row + 1, header_row + 1 + limit)
            if hasattr(values.get(f"A{row}"), "year")]
    return pd.DataFrame(rows)


# ---------------------------------------------------------------------------
# translating the formula
# ---------------------------------------------------------------------------

def _translate(formula, values, input_letter, input_row) -> tuple:
    """An Excel formula as a Python expression in ``x``, plus the coefficients used.

    Absolute references become their cached values, the input column's own-row
    reference becomes ``x``, and Excel's ``^`` and ``EXP`` become Python's. The
    result is checked against an allowlist of characters before it is ever
    evaluated.
    """
    expression = formula.lstrip("=")
    coefficients = {}
    for coordinate in dict.fromkeys(f"{letter}{row}" for letter, row
                                    in re.findall(r"\$([A-Z]{1,2})\$(\d+)",
                                                  expression)):
        letter, row = _split(coordinate)
        value = values.get(coordinate)
        if not isinstance(value, (int, float)):
            return "", {}, f"coefficient cell {coordinate} is empty"
        label = values.get(f"{letter}{row - 1}")
        coefficients[str(label).strip() if label is not None else coordinate] = \
            float(value)
        expression = expression.replace(f"${letter}${row}", repr(float(value)))

    if input_letter:
        expression = re.sub(rf"\b\$?{input_letter}\$?{input_row}\b", "x", expression)
    expression = expression.replace("^", "**")
    expression = re.sub(r"\bEXP\(", "exp(", expression, flags=re.I)

    residue = expression.replace("exp(", "")
    if not _SAFE_EXPRESSION.match(residue):
        return "", coefficients, f"formula uses unsupported syntax: {formula}"
    return expression, coefficients, ""


def _verify(rating, visits, input_letter, estimated_letter) -> bool:
    """Does re-evaluating the expression reproduce the workbook's own column?"""
    if visits.empty or estimated_letter is None:
        return False
    inputs = pd.to_numeric(visits.get(input_letter), errors="coerce")
    expected = pd.to_numeric(visits.get(estimated_letter), errors="coerce")
    usable = inputs.notna() & expected.notna()
    if not usable.any():
        return False
    computed = rating.discharge(inputs[usable].to_numpy())
    return bool(np.allclose(computed, expected[usable].to_numpy(),
                            rtol=settings.WORKBOOK_EQUATION_CHECK_TOLERANCE, atol=1e-9))


def datum_tie(visits, rating_curve, input_letter) -> tuple:
    """The elevation tying a workbook's stage axis to its distance axis.

    Measured per visit as ``stage + distance / 304.8`` and averaged, with the spread
    across visits returned so an inconsistent tie can be refused.

    Returns
    -------
    tuple
        ``(elevation in NAVD88 feet or None, spread in feet or None)``.
    """
    if visits.empty or rating_curve is None or rating_curve.empty:
        return None, None
    merged = (rating_curve.assign(date=pd.to_datetime(rating_curve["date"]).dt.normalize())
              .merge(visits, on="date", how="inner"))
    distances = pd.to_numeric(merged.get(input_letter), errors="coerce")
    usable = distances.notna() & merged["stage_ft"].notna()
    if not usable.any():
        return None, None
    implied = merged.loc[usable, "stage_ft"] + distances[usable] / settings.MM_PER_FOOT
    return float(implied.mean()), float(implied.max() - implied.min())


# ---------------------------------------------------------------------------
# the public reads
# ---------------------------------------------------------------------------

def rating_curve_sheet(path, sheet=None) -> "pd.DataFrame | None":
    """The paired measurements behind a workbook's chart, as date / stage / discharge.

    The same contract as
    :func:`limnotech_rating_curves.data.magl.rating_curve_sheet`, read through
    openpyxl so the workbook's huge sheets are never parsed. Stage is a NAVD88
    water-surface elevation, which is what the sheet plots.

    Parameters
    ----------
    path : path-like
        The workbook.
    sheet : str, optional
        Sheet name. Found case-insensitively if not given - SSL-03 spells it
        ``rating curve``.

    Returns
    -------
    pandas.DataFrame or None
        None when the workbook has no such sheet or the sheet has no header.
    """
    import openpyxl
    path = Path(path)
    names = _sheet_names(path)
    if sheet is None:
        sheet = next((name for name in names
                      if name.strip().lower() == "rating curve"), None)
    if sheet is None or sheet not in names:
        return None
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    rows = list(book[sheet].iter_rows(min_row=1, max_row=200, max_col=6,
                                      values_only=True))
    book.close()

    start = next((index for index, row in enumerate(rows)
                  if any(isinstance(cell, str) and cell.strip().lower() == "date"
                         for cell in row)), None)
    if start is None:
        return None
    header = [str(cell).strip().lower() if cell is not None else ""
              for cell in rows[start]]
    wanted = {"date": "date", "stage": "stage_ft", "discharge": "discharge_cfs"}
    if any(name not in header for name in wanted):
        return None
    frame = pd.DataFrame(rows[start + 1:], columns=header)
    frame = frame.loc[:, ~pd.Index(header).duplicated()]
    tidy = frame[list(wanted)].rename(columns=wanted)
    for column in ("stage_ft", "discharge_cfs"):
        tidy[column] = pd.to_numeric(tidy[column], errors="coerce")
    tidy["date"] = pd.to_datetime(tidy["date"], errors="coerce")
    tidy = tidy.dropna(subset=["stage_ft", "discharge_cfs"])
    return tidy.sort_values("stage_ft").reset_index(drop=True)


def stored_equation(sensor, path=None) -> WorkbookRating:
    """The rating a workbook already carries, ready to evaluate against stage.

    Parameters
    ----------
    sensor : str
        Sensor id, e.g. ``"SBR-09"``.
    path : path-like, optional
        The workbook. Looked up from :func:`workbook_paths` if not given.

    Returns
    -------
    WorkbookRating
        With ``note`` set and ``ok`` False when the workbook has the coefficient
        layout but never had coefficients filled in - which is the case at 17 of the
        52 sensors, so it is an expected outcome rather than an error.
    """
    sensor = str(sensor).replace("_WL", "")
    if path is None:
        path = workbook_paths().get(sensor)
    if path is None:
        return WorkbookRating(sensor=sensor, path=Path(),
                              note=f"no flow workbook for {sensor}")
    path = Path(path)

    names = _sheet_names(path)
    summary = summary_sheet_name(path, names)
    if summary is None:
        return WorkbookRating(sensor=sensor, path=path,
                              note="no summary sheet with a control-point banner")

    formulas = _grid(path, summary, settings.WORKBOOK_SUMMARY_ROWS, data_only=False)
    values = _grid(path, summary, settings.WORKBOOK_SUMMARY_ROWS, data_only=True)
    header_row, headers = _header(formulas)
    if header_row is None:
        return WorkbookRating(sensor=sensor, path=path,
                              note="summary sheet has no Date/Time header row")

    estimated = _column_for(headers, "estimated")
    row, formula = _first_formula(formulas, estimated, header_row)
    if formula is None:
        return WorkbookRating(sensor=sensor, path=path,
                              note="no Estimated Discharge formula")

    inputs = [letter for absolute, letter, at_row
              in re.findall(r"(\$?)([A-Z]{1,2})\$?(\d+)", formula)
              if int(at_row) == row and not absolute]
    input_letter = inputs[0] if inputs else None
    expression, coefficients, problem = _translate(formula, values, input_letter, row)
    form = ("exponential" if "exp(" in expression
            else "quadratic" if "**2" in expression
            else "linear" if input_letter else "constant")
    rating = WorkbookRating(sensor=sensor, path=path, form=form, formula=formula,
                            expression=expression, coefficients=coefficients,
                            note=problem)
    if problem:
        return rating

    header = headers.get(input_letter, "")
    rating.variable = ("stage_ft" if "water elevation" in header.lower()
                       else "distance_mm")

    visits = _visits(values, headers, header_row)
    rating.verified = _verify(rating, visits, input_letter, estimated)
    if rating.variable == "distance_mm":
        elevation, spread = datum_tie(visits, rating_curve_sheet(path), input_letter)
        rating.radar_elevation_ft, rating.datum_spread_ft = elevation, spread
        if elevation is None:
            rating.note = ("no visit ties the stage axis to the distance axis, so "
                           "the equation cannot be drawn against stage")
        elif spread > settings.WORKBOOK_DATUM_TIE_TOLERANCE_FT:
            log.warning("%s: the stage-to-distance tie varies by %.2f ft across "
                        "visits; the drawn curve assumes the mean, %.2f ft",
                        sensor, spread, elevation)
    return rating


def reference(sensor, sample, path=None) -> "dict | None":
    """A workbook's rating packaged the way the map draws a reference curve.

    Scored on the same measurements the fitted models saw, so it lands in the
    comparison table as one more row.

    Parameters
    ----------
    sensor : str
        Sensor id.
    sample : Sample
        The measurements. Its stage must be a NAVD88 elevation, which is the axis
        the workbook's own Rating Curve sheet uses.
    path : path-like, optional
        The workbook.

    Returns
    -------
    dict or None
        ``{"label", "curve", "metrics", "rating"}``, or None when the workbook has
        no usable stored equation.
    """
    rating = stored_equation(sensor, path)
    if not rating.ok or len(sample) == 0:
        return None
    curve = rating.curve(sample.stage_ft)
    if curve.empty:
        return None
    predicted = rating.discharge_at_stage(sample.stage_ft)
    return {"label": f"spreadsheet {rating.form} ({sensor})",
            "curve": curve[["stage_ft", "discharge_cfs"]],
            "metrics": fit_metrics(sample.discharge_cfs, predicted),
            "rating": rating}


def status(root=None) -> pd.DataFrame:
    """One row per workbook describing what it holds and whether it is usable.

    Parameters
    ----------
    root : path-like, optional
        The ``magl_curve_data`` tree.

    Returns
    -------
    pandas.DataFrame
        Indexed by sensor, with the measurement count, the stored equation's form,
        axis, coefficients and verification, the datum tie, and ``fittable`` - an
        equation that evaluates plus at least three measurements.
    """
    rows = []
    for sensor, path in workbook_paths(root).items():
        measurements = rating_curve_sheet(path)
        rating = stored_equation(sensor, path)
        rows.append({
            "sensor": sensor,
            "cluster": sensor.split("-")[0],
            "file": str(Path(path).relative_to(
                Path(settings.MAGL_WORKBOOK_ROOT if root is None else root))),
            "n_points": 0 if measurements is None else len(measurements),
            "stage_min": None if measurements is None or measurements.empty
                         else round(float(measurements["stage_ft"].min()), 2),
            "stage_max": None if measurements is None or measurements.empty
                         else round(float(measurements["stage_ft"].max()), 2),
            "has_equation": rating.ok,
            "form": rating.form,
            "axis": rating.variable,
            "equation": rating.equation() if rating.ok else "",
            "verified": rating.verified,
            "radar_elevation_ft": rating.radar_elevation_ft,
            "datum_spread_ft": rating.datum_spread_ft,
            "note": rating.note,
        })
    table = pd.DataFrame(rows).set_index("sensor")
    table["fittable"] = table["has_equation"] & (table["n_points"] >= 3)
    return table


def fittable_sensors(root=None, min_points: int = 3) -> list:
    """Sensors with a usable stored equation and enough measurements to refit.

    This is the selection the comparison runs on: fitting a model at a sensor whose
    workbook has no equation would leave nothing to compare against.

    Parameters
    ----------
    root : path-like, optional
        The ``magl_curve_data`` tree.
    min_points : int, default 3
        Fewest measurements. Three is the floor for a single-segment power law.

    Returns
    -------
    list of str
    """
    table = status(root)
    return list(table[table["has_equation"] & (table["n_points"] >= min_points)].index)
